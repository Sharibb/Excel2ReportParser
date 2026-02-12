"""Check what placeholders exist in template vulnerability tables"""
from docx import Document

doc = Document("/app/test_template.docx")

print("=" * 80)
print("CHECKING TEMPLATE FOR PLACEHOLDERS")
print("=" * 80)

# Find vulnerability table templates (skip summary table)
for table_idx, table in enumerate(doc.tables):
    # Look for tables with placeholders
    table_text = ""
    for row in table.rows:
        for cell in row.cells:
            table_text += cell.text + " "
    
    # Check if this is a vulnerability table template
    if any(p in table_text for p in ['{{RISK', '{{CVSS', '{{CWE', '{{DESCRIPTION}}']):
        print(f"\nVULNERABILITY TABLE TEMPLATE #{table_idx}")
        print("-" * 80)
        print(f"Rows: {len(table.rows)}, Columns: {len(table.columns)}\n")
        
        # Show each row with its placeholders
        for row_idx, row in enumerate(table.rows):
            if len(row.cells) >= 2:
                label = row.cells[0].text.strip()
                value = row.cells[1].text.strip()
                
                # Highlight rows with or without placeholders
                if '{{' in value:
                    print(f"  Row {row_idx:2}: {label:30} | ✅ {value[:60]}")
                elif value and len(value) > 5:  # Has actual text (not placeholder)
                    print(f"  Row {row_idx:2}: {label:30} | ❌ HARDCODED: {value[:60]}")
                else:
                    print(f"  Row {row_idx:2}: {label:30} | {value}")

print("\n" + "=" * 80)
